"""Build an editable PowerPoint client report from the dashboard's report JSON.

Usage: python3 report_pptx.py <report.json> <out.pptx>      (requires: pip install python-pptx)
The dashboard posts this JSON to /api/pptx; server.js runs this script and returns the file.
"""
import json, sys, io, base64
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION

INK = RGBColor(0x15, 0x20, 0x2B); MUTED = RGBColor(0x59, 0x67, 0x78); ACCENT = RGBColor(0x1F, 0x5E, 0x8C)
SOFT = RGBColor(0xDD, 0xE9, 0xF3); LINE = RGBColor(0xD5, 0xDD, 0xE6); ROW = RGBColor(0xF3, 0xF6, 0xF9); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = 'Calibri'
W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.55)

def text(slide, x, y, w, h, s, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = s; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = FONT
    return tb

def header(slide, title, meta):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.12)); bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()
    text(slide, MARGIN, Inches(0.35), W - 2 * MARGIN, Inches(0.6), title, 24, True)
    text(slide, MARGIN, Inches(0.9), W - 2 * MARGIN, Inches(0.35), meta, 11, False, MUTED)

def footer(slide, s, n):
    text(slide, MARGIN, H - Inches(0.45), Inches(9), Inches(0.3), s, 9, False, MUTED)
    text(slide, W - MARGIN - Inches(1), H - Inches(0.45), Inches(1), Inches(0.3), str(n), 9, False, MUTED, PP_ALIGN.RIGHT)

def table(slide, headers, rows, top, row_h=0.36):
    cols = len(headers); avail = W - 2 * MARGIN
    first = min(Inches(3.6), max(Inches(1.8), avail * (0.32 if cols > 6 else 0.4)))
    other = int((avail - first) / max(1, cols - 1))
    height = Inches(row_h) * (len(rows) + 1)
    shp = slide.shapes.add_table(len(rows) + 1, cols, MARGIN, top, avail, height); t = shp.table
    t.columns[0].width = int(first)
    for j in range(1, cols): t.columns[j].width = other
    size = 11 if cols <= 6 else 10 if cols <= 9 else 9
    def cell(i, j, val, head=False, total=False):
        c = t.cell(i, j); c.text = '' ; tf = c.text_frame; tf.word_wrap = True
        c.margin_left = c.margin_right = Inches(0.06); c.margin_top = c.margin_bottom = Inches(0.03)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
        r = p.add_run(); r.text = str(val); r.font.size = Pt(size); r.font.name = FONT
        r.font.bold = head or total; r.font.color.rgb = WHITE if head else INK
        c.fill.solid(); c.fill.fore_color.rgb = ACCENT if head else (SOFT if total else (ROW if i % 2 == 0 else WHITE))
    for j, h in enumerate(headers): cell(0, j, h, head=True)
    for i, r in enumerate(rows, 1):
        tot = str(r[0]).lower().startswith('total') or str(r[0]).startswith(('Budget for media', 'Remaining'))
        for j in range(cols): cell(i, j, r[j] if j < len(r) else '', total=tot)

def build(d, out):
    prs = Presentation(); prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]; n = 0
    meta = f"{d['brand']} · {d['project']} · {d['period']}"
    # title
    s = prs.slides.add_slide(blank); n += 1
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H); bg.fill.solid(); bg.fill.fore_color.rgb = ACCENT; bg.line.fill.background()
    text(s, MARGIN, Inches(2.2), W - 2 * MARGIN, Inches(1.1), f"{d['brand']} – {d['project']}", 40, True, WHITE)
    text(s, MARGIN, Inches(3.35), W - 2 * MARGIN, Inches(0.6), f"{d['title']} · {d['period']}", 20, False, WHITE)
    text(s, MARGIN, Inches(4.1), W - 2 * MARGIN, Inches(0.5), f"{d.get('location', '')} · Prepared {d['prepared']}", 13, False, SOFT)
    # KPI summary
    s = prs.slides.add_slide(blank); n += 1; header(s, 'Summary', meta)
    k = d['kpis']; gap = Inches(0.25); bw = int((W - 2 * MARGIN - gap * (len(k) - 1)) / len(k))
    for i, (label, value, sub) in enumerate(k):
        x = MARGIN + i * (bw + gap); box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.45), bw, Inches(1.4))
        box.fill.solid(); box.fill.fore_color.rgb = ROW; box.line.color.rgb = LINE; box.adjustments[0] = 0.08
        text(s, x + Inches(0.2), Inches(1.57), bw - Inches(0.4), Inches(0.3), label.upper(), 10, True, MUTED)
        text(s, x + Inches(0.2), Inches(1.85), bw - Inches(0.4), Inches(0.6), value, 26, True, INK)
        text(s, x + Inches(0.2), Inches(2.45), bw - Inches(0.4), Inches(0.3), sub, 11, False, MUTED)
    first_table = next((x for x in d['sections'] if x['type'] == 'table'), None)
    if first_table and len(first_table['rows']) <= 12:
        table(s, first_table['headers'], first_table['rows'], Inches(3.1), row_h=0.29)
    else:
        first_table = None
    footer(s, d.get('footer', ''), n)
    # sections
    for sec in d['sections']:
        if sec is first_table: continue
        if sec['type'] == 'table':
            rows = sec['rows']; pages = max(1, -(-len(rows) // 13)); per = -(-len(rows) // pages) if rows else 1
            chunks = [rows[i:i + per] for i in range(0, len(rows), per)] or [[]]
            for ci, ch in enumerate(chunks):
                s = prs.slides.add_slide(blank); n += 1
                header(s, sec['title'] + (f' ({ci + 1}/{len(chunks)})' if len(chunks) > 1 else ''), meta)
                table(s, sec['headers'], ch, Inches(1.4)); footer(s, d.get('footer', ''), n)
        elif sec['type'] == 'pie':
            s = prs.slides.add_slide(blank); n += 1; header(s, sec['title'], meta)
            series = [x for x in sec.get('series', []) if any(p[1] for p in x['points'])]
            cw = Inches(4.0)
            for i, sr in enumerate(series):
                pts = [p for p in sr['points'] if p[1]]
                cd = CategoryChartData(); cd.categories = [p[0] for p in pts]; cd.add_series(sr['name'], [p[1] for p in pts])
                gf = s.shapes.add_chart(XL_CHART_TYPE.PIE, MARGIN + i * (cw + Inches(0.2)), Inches(1.35), cw, Inches(3.3), cd); ch = gf.chart
                ch.has_title = True; ch.chart_title.text_frame.text = sr['name']; ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
                ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(10)
                pl = ch.plots[0]; pl.has_data_labels = True; dl = pl.data_labels; dl.number_format = '0%'; dl.show_percentage = True; dl.show_value = False; dl.show_category_name = True; dl.font.size = Pt(11); dl.font.bold = True; dl.font.color.rgb = WHITE
                dl.position = XL_LABEL_POSITION.BEST_FIT
                pcol = {'Meta': ACCENT, 'Google': RGBColor(0xC8, 0x81, 0x1A), 'LinkedIn': RGBColor(0x5B, 0x7F, 0x3A), 'Taboola': RGBColor(0x8E, 0x5A, 0x9E)}
                cols = [ACCENT, RGBColor(0xC8, 0x81, 0x1A), RGBColor(0x5B, 0x7F, 0x3A), RGBColor(0x8E, 0x5A, 0x9E), MUTED, LINE]
                for j, pt in enumerate(pl.series[0].points):
                    pt.format.fill.solid(); pt.format.fill.fore_color.rgb = pcol.get(pts[j][0], cols[j % len(cols)]); pt.format.line.color.rgb = WHITE
            t = sec.get('table')
            if t:
                x0 = MARGIN + len(series) * (cw + Inches(0.2)); w = W - MARGIN - x0
                if w < Inches(4): table(s, t['headers'], t['rows'], Inches(4.9), row_h=0.3)
                else:
                    shp = s.shapes.add_table(len(t['rows']) + 1, len(t['headers']), x0, Inches(1.5), w, Inches(0.33) * (len(t['rows']) + 1)); tb = shp.table
                    for j, h in enumerate(t['headers']):
                        c = tb.cell(0, j); c.text = h; c.fill.solid(); c.fill.fore_color.rgb = ACCENT; r = c.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(10); r.font.bold = True; r.font.color.rgb = WHITE
                    for i2, row in enumerate(t['rows'], 1):
                        for j, v in enumerate(row):
                            c = tb.cell(i2, j); c.text = str(v); c.fill.solid(); c.fill.fore_color.rgb = ROW if i2 % 2 == 0 else WHITE
                            p = c.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT; p.runs[0].font.size = Pt(10)
            footer(s, d.get('footer', ''), n)
        elif sec['type'] == 'cards':
            items = sec['items']; per = 5
            for ci in range(0, len(items), per):
                ch = items[ci:ci + per]; s = prs.slides.add_slide(blank); n += 1
                pages = -(-len(items) // per)
                header(s, sec['title'] + (f' ({ci // per + 1}/{pages})' if pages > 1 else ''), meta)
                cw = int((W - 2 * MARGIN - Inches(0.25) * (per - 1)) / per)
                for i, it in enumerate(ch):
                    x = MARGIN + i * (cw + Inches(0.25)); y = Inches(1.4)
                    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, cw, Inches(5.5)); box.fill.solid(); box.fill.fore_color.rgb = ROW; box.line.color.rgb = LINE
                    ih = Inches(2.6)
                    if it.get('thumb'):
                        try:
                            blob = base64.b64decode(it['thumb'].split(',', 1)[1]); pic = s.shapes.add_picture(io.BytesIO(blob), x + Inches(0.1), y + Inches(0.1), height=ih)
                            if pic.width > cw - Inches(0.2):
                                r = (cw - Inches(0.2)) / pic.width; pic.width = int(pic.width * r); pic.height = int(pic.height * r)
                            pic.left = int(x + (cw - pic.width) / 2)
                        except Exception: pass
                    else:
                        text(s, x, y + Inches(1.1), cw, Inches(0.4), 'No image', 11, False, MUTED, PP_ALIGN.CENTER)
                    text(s, x + Inches(0.12), y + ih + Inches(0.2), cw - Inches(0.24), Inches(0.5), f"#{it['rank']}  {it['name']}", 11, True, INK)
                    tb = s.shapes.add_textbox(x + Inches(0.12), y + ih + Inches(0.75), cw - Inches(0.24), Inches(2)); tf = tb.text_frame; tf.word_wrap = True
                    for k, (lab, val) in enumerate(it['metrics']):
                        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                        r = p.add_run(); r.text = f'{lab}  '; r.font.size = Pt(10); r.font.color.rgb = MUTED; r.font.name = FONT
                        r = p.add_run(); r.text = str(val); r.font.size = Pt(10); r.font.bold = True; r.font.color.rgb = INK; r.font.name = FONT
                footer(s, d.get('footer', ''), n)
        else:
            items = sec['items']; pages = max(1, -(-len(items) // 6)); per = -(-len(items) // pages) if items else 1
            for ci in range(0, max(1, len(items)), per):
                ch = items[ci:ci + per]
                s = prs.slides.add_slide(blank); n += 1
                header(s, sec['title'] + (f' ({ci // per + 1}/{pages})' if pages > 1 else ''), meta)
                tb = s.shapes.add_textbox(MARGIN, Inches(1.45), W - 2 * MARGIN, H - Inches(2.1)); tf = tb.text_frame; tf.word_wrap = True
                for i, it in enumerate(ch):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.space_after = Pt(12)
                    r = p.add_run(); r.text = f'{ci + i + 1}.  '; r.font.bold = True; r.font.color.rgb = ACCENT; r.font.size = Pt(16); r.font.name = FONT
                    r = p.add_run(); r.text = it; r.font.size = Pt(16); r.font.color.rgb = INK; r.font.name = FONT
                footer(s, d.get('footer', ''), n)
    s = prs.slides.add_slide(blank); n += 1
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H); bg.fill.solid(); bg.fill.fore_color.rgb = ACCENT; bg.line.fill.background()
    text(s, MARGIN, Inches(3.1), W - 2 * MARGIN, Inches(1), 'Thank you', 40, True, WHITE, PP_ALIGN.CENTER)
    prs.save(out)

if __name__ == '__main__':
    build(json.load(open(sys.argv[1], encoding='utf-8')), sys.argv[2])
